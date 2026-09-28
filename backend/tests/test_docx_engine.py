import io

from docx import Document

from app.engine import docx_engine
from app.engine.docx_engine import render_resume


def _composition(layout=None):
    return {
        "personal": {"name": "Ada Lovelace", "email": "ada@example.com"},
        "summary": "Engineer.",
        "experience": [{"title": "Engineer", "company": "Acme", "start_date": "2020",
                        "end_date": "Present", "bullets": ["Built things"]}],
        "section_order": ["summary", "experience"],
        "layout": layout or {},
    }


def test_layout_override_does_not_leak_into_shared_theme():
    before = docx_engine._THEMES["ats"].margin_inches
    render_resume(_composition({"margin_inches": 1.0, "font_size_body": 11}), "ats")
    assert docx_engine._THEMES["ats"].margin_inches == before
    assert docx_engine._THEMES["ats"].font_size_body == 10


def test_layout_values_are_clamped():
    data = render_resume(_composition({"margin_inches": 0.01, "font_size_body": 40}), "ats")
    doc = Document(io.BytesIO(data))
    assert abs(doc.sections[0].left_margin.inches - 0.4) < 0.01


def test_garbage_layout_values_are_ignored():
    data = render_resume(_composition({"margin_inches": "wide", "font_size_body": True}), "swe")
    assert Document(io.BytesIO(data)).paragraphs  # renders a valid document


def test_unknown_format_falls_back_to_ats():
    data = render_resume(_composition(), "not-a-format")
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "ADA LOVELACE" in text
