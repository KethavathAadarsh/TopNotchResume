"""
File upload endpoint — accepts PDF/DOCX/TXT of an existing resume
and extracts text for the frontend to pre-populate the form.

File type is resolved from the filename extension first (browsers are
inconsistent about MIME types — .docx frequently arrives as
application/octet-stream), with the declared content type as a fallback.
"""
import io
import logging
import os

from fastapi import APIRouter, UploadFile, File, HTTPException

logger = logging.getLogger(__name__)
router = APIRouter()

MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 MB

# extension -> internal kind
EXT_MAP = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "text",
    ".md": "text",
    ".markdown": "text",
    ".text": "text",
}

# content-type -> internal kind (fallback when the extension is missing/unknown)
CONTENT_TYPE_MAP = {
    "application/pdf": "pdf",
    "application/x-pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "text",
    "text/markdown": "text",
}

# Legacy / unsupported formats we can name explicitly for a better error message
REJECT_EXT = {
    ".doc": "Legacy .doc files are not supported. Re-save as .docx or PDF and try again.",
    ".pages": "Apple Pages files are not supported. Export as PDF or DOCX and try again.",
    ".rtf": "RTF files are not supported. Save as PDF, DOCX, or TXT and try again.",
    ".odt": "OpenDocument files are not supported. Export as PDF or DOCX and try again.",
}


def _resolve_kind(filename: str | None, content_type: str | None) -> str:
    ext = os.path.splitext(filename or "")[1].lower()

    if ext in REJECT_EXT:
        raise HTTPException(status_code=415, detail=REJECT_EXT[ext])

    kind = EXT_MAP.get(ext)
    if kind:
        return kind

    kind = CONTENT_TYPE_MAP.get((content_type or "").split(";")[0].strip().lower())
    if kind:
        return kind

    raise HTTPException(
        status_code=415,
        detail=f"Unsupported file type: {ext or content_type or 'unknown'}. Use PDF, DOCX, or TXT.",
    )


@router.post("/upload/resume")
async def upload_resume(file: UploadFile = File(...)):
    """
    Extract text from an uploaded resume file.
    Returns raw text for the client to parse/pre-fill the form.
    """
    kind = _resolve_kind(file.filename, file.content_type)

    contents = await file.read()

    if not contents:
        raise HTTPException(status_code=400, detail="File is empty.")

    if len(contents) > MAX_FILE_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"File is too large ({len(contents) / 1_048_576:.1f} MB). Maximum is 10 MB.",
        )

    try:
        if kind == "pdf":
            text = _extract_pdf(contents)
        elif kind == "docx":
            text = _extract_docx(contents)
        else:
            text = contents.decode("utf-8", errors="ignore")
    except HTTPException:
        raise
    except Exception as e:
        logger.error("File extraction failed for %s: %s", file.filename, e)
        raise HTTPException(status_code=422, detail=f"Could not read this file: {e}")

    text = text.strip()

    if not text:
        detail = "No text found in this file."
        if kind == "pdf":
            detail = (
                "No text found in this PDF — it looks like a scanned image. "
                "Try exporting a text-based PDF, or paste the text manually."
            )
        raise HTTPException(status_code=422, detail=detail)

    logger.info("Extracted %d chars from %s (%s)", len(text), file.filename, kind)
    return {"filename": file.filename, "text": text, "char_count": len(text)}


def _extract_pdf(data: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise HTTPException(status_code=501, detail="PDF parsing requires PyMuPDF (pip install PyMuPDF)")

    with fitz.open(stream=data, filetype="pdf") as doc:
        return "\n".join(page.get_text() for page in doc)


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]

    # Many resume templates lay content out in tables — pull those too.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                # de-dupe merged cells that repeat the same text across a row
                deduped = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]]
                parts.append(" | ".join(deduped))

    return "\n".join(parts)
